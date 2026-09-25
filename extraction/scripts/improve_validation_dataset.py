#!/usr/bin/env python3
"""Remove label leakage and diversify the local 100-email validation dataset."""
from __future__ import annotations

import csv
import mimetypes
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from email.utils import parseaddr
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

import pymupdf
import qrcode
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation

from email_pipeline.attachments import AttachmentManager
from email_pipeline.config import Settings
from email_pipeline.mail.parser import MailParser
from email_pipeline.urls.extractor import collect_urls


ROOT = Path(__file__).resolve().parents[1] / "datasets" / "phishing_validation"
EMAILS = ROOT / "emails"
MANIFEST = ROOT / "manifest.csv"

SAFE_AUTH = (
    [("pass", "pass", "pass")] * 25
    + [("fail", "none", "fail")] * 8
    + [("none", "none", "none")] * 9
    + [("pass", "fail", "pass")] * 8
)
PHISHING_AUTH = (
    [("fail", "fail", "fail")] * 15
    + [("pass", "pass", "pass")] * 20
    + [("none", "none", "none")] * 8
    + [("pass", "pass", "fail")] * 7
)


def address_domain(value: str | None) -> str:
    address = parseaddr(value or "")[1]
    return address.rsplit("@", 1)[1].lower() if "@" in address else ""


def set_header(message: EmailMessage, name: str, value: str | None) -> None:
    if name in message:
        del message[name]
    if value:
        message[name] = value


def pdf_text() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Validation invoice 4582 - total EUR 4850 - https://billing.example/pay")
    data = document.tobytes()
    document.close()
    return data


def pdf_scanned() -> bytes:
    image = Image.new("RGB", (1200, 500), "white")
    ImageDraw.Draw(image).text((80, 220), "SCANNED VALIDATION INVOICE 991", fill="black")
    png = BytesIO()
    image.save(png, format="PNG")
    document = pymupdf.open()
    page = document.new_page(width=1200, height=500)
    page.insert_image(page.rect, stream=png.getvalue())
    data = document.tobytes()
    document.close()
    return data


def qr_png() -> bytes:
    output = BytesIO()
    qrcode.make("https://qr-validation.example/login").save(output, format="PNG")
    return output.getvalue()


def text_png() -> bytes:
    image = Image.new("RGB", (1000, 300), "white")
    ImageDraw.Draw(image).text((60, 130), "VALIDATION IMAGE RECEIPT 2026", fill="black")
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def docx_file() -> bytes:
    document = Document()
    document.add_heading("Validation document", 0)
    document.add_paragraph("This harmless DOCX validates passive Office text extraction.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def xlsx_file() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invoice"
    sheet.append(["Item", "Amount"])
    sheet.append(["Validation service", 1250])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def pptx_file() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Validation presentation"
    slide.placeholders[1].text = "Harmless PPTX extraction sample"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def encrypted_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Encrypted validation document")
    data = document.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="validation-owner",
        user_pw="validation-locked",
    )
    document.close()
    return data


ATTACHMENTS = (
    ("invoice_text.pdf", "application", "pdf", pdf_text),
    ("invoice_scanned.pdf", "application", "pdf", pdf_scanned),
    ("login_qr.png", "image", "png", qr_png),
    ("receipt.png", "image", "png", text_png),
    ("document.docx", "application", "vnd.openxmlformats-officedocument.wordprocessingml.document", docx_file),
    ("table.xlsx", "application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet", xlsx_file),
    ("slides.pptx", "application", "vnd.openxmlformats-officedocument.presentationml.presentation", pptx_file),
    ("corrupt.pdf", "application", "pdf", lambda: b"%PDF-1.7\ncorrupt validation sample"),
    ("unknown.bin", "application", "octet-stream", lambda: bytes(range(256)) * 4),
    ("encrypted.pdf", "application", "pdf", encrypted_pdf),
)


def replace_attachment(message: EmailMessage, spec: tuple) -> None:
    filename, maintype, subtype, factory = spec
    parts = [
        part
        for part in message.walk()
        if not part.is_multipart()
        and (part.get_content_disposition() == "attachment" or part.get_filename())
    ]
    if not parts:
        raise ValueError("Expected an attachment part")
    part = parts[0]
    part.clear_content()
    part.set_content(factory(), maintype=maintype, subtype=subtype)
    part.add_header("Content-Disposition", "attachment", filename=filename)


def main() -> None:
    with MANIFEST.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
        fieldnames = list(rows[0])

    safe_index = phishing_index = 0
    attachment_messages: list[EmailMessage] = []
    messages: dict[str, EmailMessage] = {}
    for row in rows:
        path = EMAILS / row["file_name"]
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        label = row["label_name"]
        if label == "safe":
            spf, dkim, dmarc = SAFE_AUTH[safe_index]
            if safe_index % 3 == 0:
                set_header(message, "Reply-To", str(message.get("From", "")))
            else:
                set_header(message, "Reply-To", None)
            safe_index += 1
        else:
            spf, dkim, dmarc = PHISHING_AUTH[phishing_index]
            if phishing_index % 2 == 1:
                set_header(message, "Reply-To", None)
            phishing_index += 1

        from_domain = address_domain(message.get("From")) or "unknown.invalid"
        set_header(
            message,
            "Authentication-Results",
            f"validation.local; spf={spf}; dkim={dkim} header.d={from_domain}; "
            f"dmarc={dmarc} header.from={from_domain}",
        )
        row["spf_status"], row["dkim_status"], row["dmarc_status"] = spf, dkim, dmarc
        messages[row["mail_id"]] = message
        if any(part.get_filename() for part in message.walk()):
            attachment_messages.append(message)

    if len(attachment_messages) != len(ATTACHMENTS):
        raise ValueError(f"Expected {len(ATTACHMENTS)} attachment emails, found {len(attachment_messages)}")
    for message, attachment in zip(attachment_messages, ATTACHMENTS, strict=True):
        replace_attachment(message, attachment)

    for row in rows:
        (EMAILS / row["file_name"]).write_bytes(messages[row["mail_id"]].as_bytes(policy=policy.SMTP))

    parser = MailParser()
    with TemporaryDirectory(prefix="email-validation-") as temporary:
        manager = AttachmentManager(
            Settings(data_dir=Path(temporary), ocr_lang="eng", extractor_timeout_seconds=20)
        )
        for row in rows:
            path = EMAILS / row["file_name"]
            parsed = parser.parse(path.read_bytes())
            extracted = manager.process_all(row["mail_id"], parsed.attachments)
            urls = collect_urls(
                parsed.body_text,
                parsed.body_html,
                extracted,
                parsed.body_text_derived_from_html,
            )
            row["subject"] = parsed.subject
            row["sender"] = str(messages[row["mail_id"]].get("From", ""))
            row["sender_domain"] = address_domain(parsed.from_address)
            row["receiver"] = ";".join(parsed.to)
            row["receiver_domain"] = ";".join(
                dict.fromkeys(filter(None, (address_domain(value) for value in parsed.to)))
            )
            row["reply_to"] = parsed.reply_to_address or ""
            row["date"] = str(messages[row["mail_id"]].get("Date", ""))
            row["url_count"] = str(len(urls))
            row["has_attachments"] = str(bool(parsed.attachments))
            row["attachment_name"] = ";".join(item.filename for item in parsed.attachments)
            row["attachment_type"] = ";".join(item.content_type for item in parsed.attachments)

    with MANIFEST.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    print("Improved 100 emails and regenerated manifest.csv")


if __name__ == "__main__":
    main()

