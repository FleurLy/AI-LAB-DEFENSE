from __future__ import annotations

from io import BytesIO

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


def extract_docx(payload: bytes) -> str:
    document = Document(BytesIO(payload))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(line for line in blocks if line.strip())


def extract_xlsx(payload: bytes) -> str:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True, keep_links=False)
    try:
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"[{sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    lines.append("\t".join(values))
        return "\n".join(lines)
    finally:
        workbook.close()


def extract_pptx(payload: bytes) -> str:
    presentation = Presentation(BytesIO(payload))
    lines: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines.append(f"[Slide {index}]")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                lines.append(shape.text)
    return "\n".join(lines)

